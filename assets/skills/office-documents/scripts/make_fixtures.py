#!/usr/bin/env python
"""Generate small documents that exercise the awkward parts of the format.

Reviewing comments, tracked changes or watermarks is impossible without files
that contain them, and hand-building those by clicking in Word is not repeatable.
Each fixture targets one feature so a script's report can be judged against a
known-clean or known-broken input.

Exit codes: 0 success, 2 bad input.
"""

from __future__ import annotations

import argparse
import sys
import tempfile
from pathlib import Path

from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx_kit as kit

from docx import Document
from docx.shared import Inches


def base_document(title: str, paragraphs: list[str]) -> Document:
    document = Document()
    document.add_heading(title, level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    return document


def mark_paragraph(paragraph, before: str, deleted: str, inserted: str, after: str, author: str, date: str) -> None:
    rpr = kit.first_rpr(paragraph)
    kit.clear_paragraph_content(paragraph)
    paragraph.append(kit.make_run(before, etree.fromstring(etree.tostring(rpr)) if rpr is not None else None))
    deletion = etree.SubElement(paragraph, kit.q("del"))
    deletion.set(kit.q("id"), "101")
    deletion.set(kit.q("author"), author)
    deletion.set(kit.q("date"), date)
    deletion.append(kit.make_run(deleted, tag="delText"))
    insertion = etree.SubElement(paragraph, kit.q("ins"))
    insertion.set(kit.q("id"), "102")
    insertion.set(kit.q("author"), author)
    insertion.set(kit.q("date"), date)
    insertion.append(kit.make_run(inserted))
    paragraph.append(kit.make_run(after))


def make_tracked(target: Path) -> dict:
    document = base_document("Revisions demo", ["Regional revenue grew by 12% in the quarter.", "Hardware sales stayed flat."])
    document.save(target)
    doc = kit.Docx(target)
    paragraphs = doc.paragraphs()
    mark_paragraph(paragraphs[1], "Regional revenue grew by ", "12%", "14%", " in the quarter.", "Reviewer", "2024-05-01T09:00:00Z")
    mark_paragraph(paragraphs[2], "Hardware sales ", "stayed", "remained", " flat.", "Reviewer", "2024-05-01T09:01:00Z")
    settings = doc.xml("word/settings.xml")
    if settings.find(kit.q("trackRevisions")) is None:
        settings.append(settings.makeelement(kit.q("trackRevisions"), {}))
    doc.save(target)
    return {"path": str(target), "tracked_changes": 4}


def make_watermark(target: Path) -> dict:
    document = base_document("Watermark demo", ["This copy is for internal review."])
    document.save(target)
    doc = kit.Docx(target)
    part = doc.ensure_header_footer("header")
    doc.xml(part).append(kit.watermark_paragraph("DRAFT", doc.next_id()))
    doc.save(target)
    return {"path": str(target), "watermark": "DRAFT"}


def make_comments(target: Path) -> dict:
    document = base_document("Comments demo", ["The first paragraph carries a question.", "The second paragraph is clean."])
    document.save(target)
    doc = kit.Docx(target)
    root = etree.Element(kit.q("comments"), nsmap={"w": kit.W, "r": kit.R})
    doc.set_xml("word/comments.xml", root)
    doc.ensure_override("word/comments.xml", kit.CT_BY_REL[kit.REL_COMMENTS])
    doc.ensure_rel("word/document.xml", kit.REL_COMMENTS, "comments.xml")
    comment = etree.SubElement(root, kit.q("comment"))
    comment.set(kit.q("id"), "1")
    comment.set(kit.q("author"), "Reviewer")
    comment.set(kit.q("initials"), "RV")
    comment.set(kit.q("date"), "2024-05-02T10:00:00Z")
    body = etree.SubElement(comment, kit.q("p"))
    etree.SubElement(etree.SubElement(body, kit.q("r")), kit.q("t")).text = "Is this number current?"

    paragraph = doc.paragraphs()[1]
    start = etree.Element(kit.q("commentRangeStart"))
    start.set(kit.q("id"), "1")
    paragraph.insert(1, start)
    end = etree.SubElement(paragraph, kit.q("commentRangeEnd"))
    end.set(kit.q("id"), "1")
    run = etree.SubElement(paragraph, kit.q("r"))
    etree.SubElement(run, kit.q("commentReference")).set(kit.q("id"), "1")
    doc.save(target)
    return {"path": str(target), "comments": 1}


def make_edge_cases(target: Path, image_dir: Path) -> dict:
    image_path = image_dir / "sample.png"
    try:
        from PIL import Image

        Image.new("RGB", (320, 200), "#4472c4").save(image_path)
    except ImportError:
        image_path = None

    document = Document()
    document.add_heading("Edge cases", level=1)
    document.add_paragraph("")
    document.add_heading("A level three heading after a level one", level=3)
    document.add_paragraph("A very long paragraph " * 40)
    document.add_paragraph("Contact: writer@example.org or +1 555 010 9999.")
    table = document.add_table(rows=2, cols=10)
    for row_index in range(2):
        for column_index in range(10):
            table.cell(row_index, column_index).text = "R%dC%d" % (row_index + 1, column_index + 1)
    if image_path is not None:
        document.add_picture(str(image_path), width=Inches(4))
    document.save(target)
    return {"path": str(target), "fixtures": ["empty_paragraph", "skipped_heading_level", "wide_table", "image_without_alt", "contact_details"]}


BUILDERS = {
    "tracked": make_tracked,
    "watermark": make_watermark,
    "comments": make_comments,
}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Create sample .docx files for tracked changes, watermarks, comments and edge cases.")
    parser.add_argument("--outdir", required=True, help="Directory for the generated files")
    parser.add_argument("--only", default="all", help="tracked, watermark, comments, edge, or all")
    args = parser.parse_args(argv)

    out_dir = Path(args.outdir).expanduser().resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    wanted = [name.strip() for name in args.only.split(",") if name.strip()]
    if "all" in wanted:
        wanted = ["tracked", "watermark", "comments", "edge"]
    unknown = [name for name in wanted if name not in BUILDERS and name != "edge"]
    if unknown:
        print("error: unknown fixture %s" % ", ".join(unknown), file=sys.stderr)
        return 2

    results = []
    try:
        for name in wanted:
            if name == "edge":
                results.append(make_edge_cases(out_dir / "edge_cases.docx", out_dir))
            else:
                results.append(BUILDERS[name](out_dir / ("%s.docx" % name)))
    except (OSError, KeyError) as exc:
        print("error: %s" % exc, file=sys.stderr)
        return 2

    kit.emit({"ok": True, "outdir": str(out_dir), "fixtures": results})
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
