#!/usr/bin/env python3
"""Inventory an existing .xlsx workbook before touching it.

Read-only: opens the file, never saves it. Prints one JSON object listing, per
sheet, the structures that openpyxl can silently rewrite on save (Excel tables,
pivot caches, shared and array formulas, conditional formats, data validations,
charts, merged ranges, defined names), plus whether cached formula results are
present. Use it to plan an edit from facts instead of guesses.

Exit codes:
  0  workbook inspected, JSON printed on stdout
  1  the file is missing or is not a readable .xlsx container
  2  usage error (argparse)
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from pathlib import Path

# Windows consoles default to a legacy code page; every payload is UTF-8 JSON.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

if hasattr(sys.stdout, "reconfigure"):
    # Legacy console codepages (cp936, cp1252) must not abort a report.
    sys.stdout.reconfigure(errors="replace")

FORMULA_KINDS = {
    "shared": '<f t="shared"',
    "array": '<f t="array"',
    "data_table": 't="dataTable"',
}
ATTR = re.compile(r'([\w:]+)="([^"]*)"')


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Inventory an .xlsx workbook: sheets, formulas, tables, "
        "pivot caches, conditional formats, defined names, cached values.",
    )
    parser.add_argument("--input", required=True, help="path to the .xlsx file")
    parser.add_argument(
        "--no-values",
        action="store_true",
        help="skip the second load that checks cached formula results",
    )
    parser.add_argument(
        "--sample",
        type=int,
        default=5,
        help="how many formula samples to show per sheet (default 5)",
    )
    return parser.parse_args()


# --------------------------------------------------------------------------
# Raw container scan: what the file holds that a library round-trip may drop.
# --------------------------------------------------------------------------


def _attrs(tag: str) -> dict[str, str]:
    return dict(ATTR.findall(tag))


def _sheet_xml_names(archive: zipfile.ZipFile) -> dict[str, str]:
    """Map worksheet part path -> sheet name, following workbook relationships."""
    try:
        workbook = archive.read("xl/workbook.xml").decode("utf-8", "replace")
        rels = archive.read("xl/_rels/workbook.xml.rels").decode("utf-8", "replace")
    except KeyError:
        return {}
    rel_targets: dict[str, str] = {}
    for tag in re.findall(r"<Relationship\b[^>]*>", rels):
        attrs = _attrs(tag)
        rel_targets[attrs.get("Id", "")] = attrs.get("Target", "")

    mapping: dict[str, str] = {}
    for tag in re.findall(r"<sheet\b[^>]*?/?>", workbook):
        attrs = _attrs(tag)
        target = rel_targets.get(attrs.get("r:id") or attrs.get("rId") or "", "")
        if not target:
            continue
        if target.startswith("/"):
            path = target.lstrip("/")
        else:
            path = "xl/" + target
        parts: list[str] = []
        for segment in path.split("/"):
            if segment == "..":
                if parts:
                    parts.pop()
            elif segment not in ("", "."):
                parts.append(segment)
        mapping["/".join(parts)] = attrs.get("name", "")
    return mapping


def _formula_kinds(archive: zipfile.ZipFile, part: str) -> dict[str, int]:
    try:
        xml = archive.read(part).decode("utf-8", "replace")
    except KeyError:
        return {kind: 0 for kind in FORMULA_KINDS}
    return {kind: xml.count(marker) for kind, marker in FORMULA_KINDS.items()}


def _container(archive: zipfile.ZipFile) -> dict[str, object]:
    names = archive.namelist()

    def count(prefix: str, suffix: str = "") -> int:
        return sum(1 for name in names if name.startswith(prefix) and name.endswith(suffix))

    return {
        "tables": count("xl/tables/", ".xml"),
        "pivot_tables": count("xl/pivotTables/", ".xml"),
        "pivot_caches": count("xl/pivotCache/", ".xml"),
        "charts": count("xl/charts/", ".xml"),
        "drawings": count("xl/drawings/drawing", ".xml"),
        "media": count("xl/media/"),
        "comment_parts": count("xl/comments", ".xml"),
        "external_links": count("xl/externalLinks/", ".xml"),
        "has_vba_project": "xl/vbaProject.bin" in names,
        "has_calc_chain": "xl/calcChain.xml" in names,
        "has_shared_strings": "xl/sharedStrings.xml" in names,
    }


# --------------------------------------------------------------------------
# Sheet inventory through openpyxl.
# --------------------------------------------------------------------------


def _anchor(chart) -> str | None:
    from openpyxl.utils import get_column_letter

    start = getattr(getattr(chart, "anchor", None), "_from", None)
    if start is None:
        return None
    return f"{get_column_letter(start.col + 1)}{start.row + 1}"


def _chart_kind(chart) -> str:
    """Name the chart the way a user would: a column is not a bar."""
    class_name = type(chart).__name__
    kind = class_name[: -len("Chart")].lower() if class_name.endswith("Chart") else class_name.lower()
    if kind == "bar":
        return "column" if getattr(chart, "type", "col") == "col" else "bar"
    return kind


def _date_system(workbook) -> str:
    epoch = getattr(workbook, "epoch", None)
    return "1904" if getattr(epoch, "year", 1900) == 1904 else "1900"


def _defined_names(workbook) -> list[dict[str, object]]:
    found: list[dict[str, object]] = []
    for entry in workbook.defined_names.values():
        scope: object = "workbook"
        index = entry.localSheetId
        if index is not None and index < len(workbook.sheetnames):
            scope = workbook.sheetnames[index]
        found.append(
            {
                "name": entry.name,
                "refers_to": entry.attr_text,
                "scope": scope,
                "hidden": bool(entry.hidden),
            }
        )
    return found


def _sheet_inventory(
    worksheet, raw: dict[str, int], sample: int
) -> tuple[dict[str, object], set[str]]:
    formulas: list[tuple[str, str]] = []
    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formulas.append((cell.coordinate, cell.value))

    charts = []
    for chart in getattr(worksheet, "_charts", []):
        charts.append(
            {
                "kind": _chart_kind(chart),
                "class": type(chart).__name__,
                "series": len(getattr(chart, "series", []) or []),
                "anchor": _anchor(chart),
            }
        )

    tables = []
    for table in worksheet.tables.values():
        tables.append(
            {
                "name": table.name,
                "ref": table.ref,
                "style": getattr(table.tableStyleInfo, "name", None),
                "header_rows": table.headerRowCount,
                "totals_rows": table.totalsRowCount,
            }
        )

    conditional = []
    for block in worksheet.conditional_formatting:
        rules = []
        for rule in block.rules:
            rules.append(
                {
                    "type": rule.type,
                    "operator": rule.operator,
                    "formula": list(rule.formula or []),
                    "priority": rule.priority,
                }
            )
        conditional.append({"range": str(block.sqref), "rules": rules})

    validations = []
    for validation in worksheet.data_validations.dataValidation:
        validations.append(
            {
                "type": validation.type,
                "cells": str(validation.sqref),
                "formula1": validation.formula1,
                "operator": validation.operator,
            }
        )

    widths = {
        name: dimension.width
        for name, dimension in worksheet.column_dimensions.items()
        if dimension.width
    }
    if len(widths) > 26:
        widths = dict(list(widths.items())[:26])

    hidden_rows = [
        index for index, dimension in worksheet.row_dimensions.items() if dimension.hidden
    ]
    hidden_columns = [
        name for name, dimension in worksheet.column_dimensions.items() if dimension.hidden
    ]

    return {
        "name": worksheet.title,
        "state": worksheet.sheet_state,
        "dimensions": worksheet.dimensions,
        "max_row": worksheet.max_row,
        "max_column": worksheet.max_column,
        "freeze_panes": worksheet.freeze_panes,
        "gridlines": bool(worksheet.sheet_view.showGridLines),
        "auto_filter": getattr(worksheet.auto_filter, "ref", None),
        "protection": bool(worksheet.protection.sheet),
        "print_area": str(worksheet.print_area) if worksheet.print_area else None,
        "merged": [str(item) for item in worksheet.merged_cells.ranges],
        "tables": tables,
        "pivot_tables": len(getattr(worksheet, "_pivots", []) or []),
        "charts": charts,
        "images": len(getattr(worksheet, "_images", [])),
        "conditional_formats": conditional,
        "data_validations": validations,
        "hidden_rows": hidden_rows[:20],
        "hidden_columns": hidden_columns[:20],
        "column_widths": widths,
        "formulas": {
            "count": len(formulas),
            "sample": [f"{address}: {text}" for address, text in formulas[: max(sample, 0)]],
            "shared": raw.get("shared", 0),
            "array": raw.get("array", 0),
            "data_tables": raw.get("data_table", 0),
        },
    }, {address for address, _ in formulas}


def _cached_values(path: Path, formula_cells: dict[str, set[str]]) -> dict[str, dict[str, object]]:
    """Classify stored results for the formula cells found in the formula pass."""
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=True)
    result: dict[str, dict[str, object]] = {}
    try:
        for name, coordinates in formula_cells.items():
            if name not in workbook.sheetnames:
                continue
            worksheet = workbook[name]
            cached = 0
            missing = 0
            errors = []
            for coordinate in sorted(coordinates):
                value = worksheet[coordinate].value
                if value is None:
                    missing += 1
                elif isinstance(value, str) and value.startswith("#"):
                    errors.append({"cell": coordinate, "value": value})
                else:
                    cached += 1
            result[name] = {"cached": cached, "missing": missing, "errors": errors}
    finally:
        workbook.close()
    return result


def main() -> int:
    args = parse_args()
    path = Path(args.input).expanduser()
    if not path.is_file():
        print(json.dumps({"ok": False, "error": f"file not found: {path}"}), file=sys.stderr)
        return 1
    if not zipfile.is_zipfile(path):
        print(json.dumps({"ok": False, "error": f"not a zip container: {path}"}), file=sys.stderr)
        return 1

    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - environment guard
        print(json.dumps({"ok": False, "error": f"openpyxl missing: {error}"}), file=sys.stderr)
        return 1

    try:
        with zipfile.ZipFile(path) as archive:
            part_names = _sheet_xml_names(archive)
            container = _container(archive)
            raw_by_sheet = {
                name: _formula_kinds(archive, part) for part, name in part_names.items()
            }
        workbook = load_workbook(str(path), data_only=False)
    except Exception as error:  # noqa: BLE001 - report any unreadable workbook the same way
        print(
            json.dumps({"ok": False, "error": f"cannot read workbook: {error}"}),
            file=sys.stderr,
        )
        return 1

    warnings: list[str] = []
    sheets = []
    formula_cells: dict[str, set[str]] = {}
    try:
        for worksheet in workbook.worksheets:
            inventory, coordinates = _sheet_inventory(
                worksheet, raw_by_sheet.get(worksheet.title, {}), args.sample
            )
            formula_cells[worksheet.title] = coordinates
            sheets.append(inventory)
        workbook_meta = {
            "sheets": list(workbook.sheetnames),
            "sheet_count": len(workbook.sheetnames),
            "active_sheet": workbook.active.title if workbook.active else None,
            "defined_names": _defined_names(workbook),
            "full_calc_on_load": bool(
                getattr(getattr(workbook, "calculation", None), "fullCalcOnLoad", False)
            ),
            "calc_mode": getattr(getattr(workbook, "calculation", None), "calcMode", None),
            "date_system": _date_system(workbook),
        }
    finally:
        workbook.close()

    if not args.no_values:
        values = _cached_values(path, formula_cells)
        for sheet in sheets:
            sheet["values"] = values.get(sheet["name"], {"cached": 0, "missing": 0, "errors": []})
            entry = sheet["values"]
            if entry["missing"]:  # type: ignore[index]
                warnings.append(
                    f"{sheet['name']}: {entry['missing']} formula cell(s) have no cached value"  # type: ignore[index]
                )
            for error in entry["errors"]:  # type: ignore[index]
                warnings.append(
                    f"{sheet['name']}!{error['cell']} already stores {error['value']}"
                )

    if container["pivot_tables"]:
        warnings.append(
            f"{container['pivot_tables']} pivot table(s) present: a library round-trip "
            "may drop or re-key them; verify after saving"
        )
    if container["has_calc_chain"]:
        warnings.append(
            "the file carries a calculation chain; a library round-trip may leave it stale"
        )

    payload = {
        "ok": True,
        "input": str(path),
        "size_bytes": path.stat().st_size,
        "workbook": workbook_meta,
        "container": container,
        "sheets": sheets,
        "warnings": warnings,
    }
    print(json.dumps(payload, ensure_ascii=False, indent=2, default=str))
    return 0


if __name__ == "__main__":
    sys.exit(main())
