#!/usr/bin/env python3
"""Place an image file on a worksheet at a cell anchor.

Images are embedded copies: the workbook keeps the pixels it was given, so an
image that later changes on disk does not update. Prefer a native chart or a
table when the content is data; use an image only for logos, screenshots, and
diagrams that have no cell representation.

Examples:
    add_image.py --input book.xlsx --sheet Cover --path logo.png --anchor B2
    add_image.py --input book.xlsx --sheet Cover --path shot.png --anchor B10 --width 640

Exit codes:
  0  image embedded (the JSON on stdout reports the anchor and size)
  1  refused: unknown sheet or unsupported image file
  2  usage error or missing file
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

# Windows consoles default to a legacy code page; every payload is UTF-8 JSON.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="replace")

IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".emf", ".wmf"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Embed an image on an .xlsx worksheet.")
    parser.add_argument("--input", required=True, help="source workbook")
    parser.add_argument("--out", help="destination workbook (default: <input>-edited.xlsx)")
    parser.add_argument("--sheet", required=True, help="target worksheet")
    parser.add_argument("--path", required=True, help="image file to embed")
    parser.add_argument("--anchor", default="A1", help="top-left cell the image floats over")
    parser.add_argument("--width", type=int, help="display width in pixels")
    parser.add_argument("--height", type=int, help="display height in pixels")
    return parser.parse_args()


def _error(message: str, code: int) -> int:
    print(json.dumps({"ok": False, "error": message}), file=sys.stderr)
    return code


def main() -> int:
    args = parse_args()
    source = Path(args.input).expanduser()
    image_path = Path(args.path).expanduser()
    if not source.is_file():
        return _error(f"input file not found: {source}", 2)
    if not image_path.is_file():
        return _error(f"image file not found: {image_path}", 2)
    if image_path.suffix.lower() not in IMAGE_SUFFIXES:
        return _error(f"unsupported image type: {image_path.suffix}", 1)

    try:
        from openpyxl import load_workbook
        from openpyxl.drawing.image import Image as XLImage
    except ImportError as error:  # pragma: no cover - environment guard
        return _error(f"openpyxl or Pillow missing: {error}", 2)

    workbook = load_workbook(str(source))
    if args.sheet not in workbook.sheetnames:
        return _error(f"sheet not found: {args.sheet}", 1)

    try:
        image = XLImage(str(image_path))
    except Exception as error:  # noqa: BLE001 - Pillow rejects unreadable files
        return _error(f"cannot read image: {error}", 1)

    if args.width:
        image.width = args.width
    if args.height:
        image.height = args.height
    workbook[args.sheet].add_image(image, args.anchor)

    out = Path(args.out).expanduser() if args.out else source.with_name(
        f"{source.stem}-edited{source.suffix}"
    )
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(out))
    workbook.close()

    print(
        json.dumps(
            {
                "ok": True,
                "input": str(source),
                "output": str(out),
                "sheet": args.sheet,
                "anchor": args.anchor,
                "image": str(image_path),
                "width": image.width,
                "height": image.height,
            },
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
