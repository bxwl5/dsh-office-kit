#!/usr/bin/env python3
"""Add, delete, rename, or reorder worksheets in a workbook.

Sheet names are referenced by formulas, charts, and defined names, so renaming
a sheet rewrites nothing: every reference to the old name breaks. Prefer adding
a sheet over renaming one that formulas point at.

Examples:
    sheet_ops.py --input book.xlsx --add Checks
    sheet_ops.py --input book.xlsx --rename Sheet1=Assumptions
    sheet_ops.py --input book.xlsx --delete Scratch --delete Notes
    sheet_ops.py --input book.xlsx --order "Outputs,Assumptions,Build,Inputs"

Exit codes:
  0  done (the JSON on stdout lists the resulting sheet order)
  1  refused: unknown sheet, duplicate name, or an attempt to delete the last sheet
  2  usage error, missing file, or unreadable workbook
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(errors="replace")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Add, delete, rename, or reorder worksheets in an .xlsx workbook.",
    )
    parser.add_argument("--input", required=True, help="source workbook")
    parser.add_argument("--out", help="destination workbook (default: <input>-edited.xlsx)")
    parser.add_argument("--add", action="append", default=[], metavar="NAME", help="create sheet NAME")
    parser.add_argument("--delete", action="append", default=[], metavar="NAME", help="remove sheet NAME")
    parser.add_argument("--rename", action="append", default=[], metavar="OLD=NEW", help="rename a sheet")
    parser.add_argument(
        "--order",
        metavar="A,B,C",
        help="move these sheets to the front in this order; unlisted sheets keep their order",
    )
    return parser.parse_args()


def _error(message: str, code: int) -> int:
    print(json.dumps({"ok": False, "error": message}), file=sys.stderr)
    return code


def default_out(source: Path) -> Path:
    return source.with_name(f"{source.stem}-edited{source.suffix}")


def main() -> int:
    args = parse_args()
    source = Path(args.input).expanduser()
    if not source.is_file():
        return _error(f"input file not found: {source}", 2)

    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - environment guard
        return _error(f"openpyxl missing: {error}", 2)

    try:
        workbook = load_workbook(str(source))
    except Exception as error:  # noqa: BLE001 - any unreadable file is reported the same way
        return _error(f"cannot read workbook: {error}", 2)

    actions: list[str] = []
    try:
        for name in args.add:
            if name in workbook.sheetnames:
                raise ValueError(f"sheet already exists: {name}")
            workbook.create_sheet(str(name)[:31])
            actions.append(f"added {name}")

        for pair in args.rename:
            if "=" not in pair:
                raise ValueError(f"--rename expects OLD=NEW, got: {pair}")
            old, new = pair.split("=", 1)
            if old not in workbook.sheetnames:
                raise ValueError(f"sheet not found: {old}")
            if new in workbook.sheetnames:
                raise ValueError(f"sheet already exists: {new}")
            workbook[old].title = str(new)[:31]
            actions.append(f"renamed {old} -> {new}")

        for name in args.delete:
            if name not in workbook.sheetnames:
                raise ValueError(f"sheet not found: {name}")
            if len(workbook.sheetnames) == 1:
                raise ValueError("refusing to delete the last sheet; a workbook needs at least one")
            workbook.remove(workbook[name])
            actions.append(f"deleted {name}")

        if args.order:
            wanted = [part.strip() for part in args.order.split(",") if part.strip()]
            missing = [name for name in wanted if name not in workbook.sheetnames]
            if missing:
                raise ValueError(f"unknown sheet in --order: {', '.join(missing)}")
            rest = [name for name in workbook.sheetnames if name not in wanted]
            workbook._sheets = [workbook[name] for name in wanted + rest]
            actions.append("reordered " + ", ".join(wanted + rest))
    except ValueError as error:
        return _error(str(error), 1)

    out = Path(args.out).expanduser() if args.out else default_out(source)
    out.parent.mkdir(parents=True, exist_ok=True)
    names = list(workbook.sheetnames)
    workbook.save(str(out))
    workbook.close()

    print(
        json.dumps(
            {
                "ok": True,
                "input": str(source),
                "output": str(out),
                "sheetnames": names,
                "actions": actions,
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
