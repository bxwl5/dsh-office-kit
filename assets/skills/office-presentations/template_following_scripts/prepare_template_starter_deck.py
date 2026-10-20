"""Build a starter deck out of the slides the plan keeps.

Purpose: the safe way to follow a supplied design is to keep its actual slides,
not to redraw a lookalike. This copies the template file, keeps the mapped
source slides in plan order, drops the rest, then applies the element actions
that only move text: keep, delete, fill-placeholder, rewrite and
rewrite-and-reposition.

Ceiling: slides are reused, not cloned. Two output slides may not map to the
same source slide, and 'replace' is left to the author (build the new element
with a recipe, then run check_template_fidelity). Charts are never rebuilt here.

Exit codes: 0 starter written, 1 plan rejected, 2 unreadable input.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "container_tools"))

from ooxml_helpers import box, placeholder_kind, slide_size_in  # noqa: E402

# Windows consoles default to a legacy code page; every payload is UTF-8 JSON.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

PLACEHOLDER_ALIASES = {
    "title": ("TITLE", "CENTER_TITLE"),
    "subtitle": ("SUBTITLE",),
    "body": ("BODY", "OBJECT"),
}


def _slides(prs) -> list:
    return list(prs.slides)


def _drop_slide(prs, slide) -> None:
    slide_id = None
    for element in prs.slides._sldIdLst:
        if element.rId == slide.part.relate_to(slide.part, "") if False else False:
            break
    # find the sldId whose relationship points at this slide part
    rels = prs.part.rels
    rId = None
    for key, relationship in rels.items():
        if relationship.reltype.endswith("/slide") and relationship.target_part is slide.part:
            rId = key
            break
    if rId is None:
        return
    for element in list(prs.slides._sldIdLst):
        if element.rId == rId:
            slide_id = element
            break
    if slide_id is not None:
        prs.slides._sldIdLst.remove(slide_id)
    prs.part.drop_rel(rId)


def _keep_only(prs, wanted: list[int]) -> None:
    ordered = [element for element in prs.slides._sldIdLst]
    keep_rids = []
    for position in wanted:
        keep_rids.append(ordered[position - 1].rId)
    for slide in list(prs.slides):
        current = None
        for key, relationship in prs.part.rels.items():
            if relationship.reltype.endswith("/slide") and relationship.target_part is slide.part:
                current = key
                break
        if current not in keep_rids:
            _drop_slide(prs, slide)
    # reorder what survived to match the plan
    by_rid = {element.rId: element for element in prs.slides._sldIdLst}
    for element in list(prs.slides._sldIdLst):
        prs.slides._sldIdLst.remove(element)
    for rId in keep_rids:
        prs.slides._sldIdLst.append(by_rid[rId])


def _find_shape(slide, element_id: str):
    shapes = list(slide.shapes)
    if element_id.isdigit():
        index = int(element_id) - 1
        return shapes[index] if 0 <= index < len(shapes) else None
    for shape in shapes:
        if shape.name == element_id:
            return shape
    if element_id.lower().startswith("shape-") and element_id[6:].isdigit():
        index = int(element_id[6:]) - 1
        return shapes[index] if 0 <= index < len(shapes) else None
    return None


def _find_placeholder(slide, alias: str):
    wanted = PLACEHOLDER_ALIASES.get(alias.lower())
    for shape in slide.shapes:
        if not shape.is_placeholder:
            continue
        kind = (placeholder_kind(shape) or "").upper()
        if alias.isdigit() and shape.placeholder_format.idx == int(alias):
            return shape
        if wanted and kind in wanted:
            return shape
        if alias.lower() == (placeholder_kind(shape) or "").lower():
            return shape
    return None


def _set_text(shape, text: str) -> None:
    frame = shape.text_frame
    paragraphs = list(frame.paragraphs)
    first = paragraphs[0]
    runs = list(first.runs)
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run._r.getparent().remove(run._r)
        for paragraph in paragraphs[1:]:
            paragraph._p.getparent().remove(paragraph._p)
    else:
        frame.text = text


def build_starter(template_path: str | Path, plan_path: str | Path, out_path: str | Path) -> dict:
    from pptx import Presentation

    template = Path(template_path)
    plan_file = Path(plan_path)
    if not template.is_file():
        return {"ok": False, "error": {"code": "template_missing", "message": str(template)}}
    if not plan_file.is_file():
        return {"ok": False, "error": {"code": "plan_missing", "message": str(plan_file)}}

    plan = json.loads(plan_file.read_text(encoding="utf-8"))
    entries = plan.get("outputSlides") or []
    if not entries:
        return {"ok": False, "error": {"code": "plan_empty", "message": "outputSlides is empty"}}

    sources = [int(entry.get("sourceSlide", 0)) for entry in entries]
    duplicates = sorted({value for value in sources if sources.count(value) > 1})
    if duplicates:
        return {
            "ok": False,
            "error": {
                "code": "duplicate_source_slide",
                "message": f"source slides {duplicates} are mapped more than once; this builder reuses slides instead of cloning them",
            },
        }

    prs = Presentation(str(template))
    total = len(prs.slides)
    out_of_range = [value for value in sources if not 1 <= value <= total]
    if out_of_range:
        return {"ok": False, "error": {"code": "source_out_of_range", "message": f"template has {total} slides, plan asks for {out_of_range}"}}

    slide_w, slide_h = slide_size_in(prs)
    _keep_only(prs, sources)
    kept = _slides(prs)

    applied: list[dict] = []
    for position, entry in enumerate(entries):
        slide = kept[position]
        for element in entry.get("elements", []):
            element_id = str(element.get("elementId", ""))
            action = str(element.get("action", "keep"))
            target = element.get("target") or {}
            if action == "keep":
                continue
            if action == "fill-placeholder":
                alias = str(target.get("placeholder") or "")
                shape = _find_placeholder(slide, alias) if alias else _find_shape(slide, element_id)
                if shape is None:
                    applied.append({"slide": position + 1, "element": element_id, "action": action, "status": "not_found"})
                    continue
                if target.get("text") is not None:
                    _set_text(shape, str(target["text"]))
                applied.append({"slide": position + 1, "element": shape.name, "action": action, "status": "applied"})
                continue
            shape = _find_shape(slide, element_id)
            if shape is None:
                applied.append({"slide": position + 1, "element": element_id, "action": action, "status": "not_found"})
                continue
            if action == "delete":
                shape._element.getparent().remove(shape._element)
                applied.append({"slide": position + 1, "element": element_id, "action": action, "status": "applied"})
                continue
            if action in {"rewrite", "rewrite-and-reposition"}:
                _set_text(shape, str(target.get("text", "")))
                if action == "rewrite-and-reposition" and isinstance(target.get("zone"), list):
                    x, y, width, height = (float(value) for value in target["zone"])
                    shape.left = int(x * slide_w * 914400)
                    shape.top = int(y * slide_h * 914400)
                    shape.width = int(width * slide_w * 914400)
                    shape.height = int(height * slide_h * 914400)
                applied.append({"slide": position + 1, "element": element_id, "action": action, "status": "applied"})
                continue
            applied.append({"slide": position + 1, "element": element_id, "action": action, "status": "left_for_author"})

    destination = Path(out_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(destination))

    leftovers = []
    for index, slide in enumerate(kept, start=1):
        for shape in slide.shapes:
            if shape.is_placeholder and not (shape.text_frame.text or "").strip():
                leftovers.append({"slide": index, "placeholder": shape.name, "bboxIn": [round(value, 2) for value in (box(shape) or (0, 0, 0, 0))]})

    return {
        "ok": True,
        "starter": str(destination.resolve()),
        "slides": len(kept),
        "sourceSlides": sources,
        "applied": applied,
        "emptyPlaceholders": leftovers,
        "note": "empty placeholders still show their prompt in edit view; fill or delete them before delivery",
    }


def main(argv: list[str] | None = None) -> int:
    import argparse

    parser = argparse.ArgumentParser(description="Build a starter deck from the template slides a plan keeps.")
    parser.add_argument("--template", required=True, help="path to the template .pptx")
    parser.add_argument("--plan", required=True, help="path to template-frame-map.json")
    parser.add_argument("--out", required=True, help="path of the starter .pptx to write")
    args = parser.parse_args(argv)

    report = build_starter(args.template, args.plan, args.out)
    json.dump(report, sys.stdout, ensure_ascii=False, indent=2)
    sys.stdout.write("\n")
    if report.get("ok"):
        return 0
    return 2 if report["error"]["code"] in {"template_missing", "plan_missing"} else 1


if __name__ == "__main__":
    raise SystemExit(main())
