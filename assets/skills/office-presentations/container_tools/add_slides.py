"""Append slides to an existing deck using the shared slide builder.

The slide builder lives in `python/office_cli.py` (`_pptx_add_slide`), so this
tool imports it instead of reimplementing titles, bullet bodies, tables, charts
and pictures. Anything `pptx create --spec` accepts inside `slides[]` is
accepted here.

Usage
-----
    <python> add_slides.py --input deck.pptx --spec extra-slides.json
    <python> add_slides.py --input deck.pptx --slide '{"layout":"bullets","title":"Next"}' --out deck-2.pptx

Exit codes
----------
0  slides appended
2  bad arguments, invalid spec, or a slide that the builder rejected
3  the output file could not be written

The JSON report on stdout is `{ok, input, output, added, slides, counts}`.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

# Windows consoles default to a legacy code page; every payload is UTF-8 JSON.
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding="utf-8", errors="replace")
    except (AttributeError, ValueError):
        pass

PACKAGE_ROOT = Path(__file__).resolve().parents[4]
CLI_DIR = PACKAGE_ROOT / "python"
COUNTERS = ("slides", "bullets", "tables", "charts", "images")


def _emit(payload: dict, code: int = 0) -> int:
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return code


def _load_slides(spec_path: str | None, inline: list[str]) -> list[dict]:
    slides: list[dict] = []
    if spec_path:
        data = json.loads(Path(spec_path).read_text(encoding="utf-8"))
        if isinstance(data, dict):
            slides.extend(data.get("slides") or [])
        elif isinstance(data, list):
            slides.extend(data)
        else:
            raise ValueError("spec must be an object with 'slides' or a list of slides")
    for raw in inline:
        slides.append(json.loads(raw))
    for slide in slides:
        if not isinstance(slide, dict):
            raise ValueError(f"slide must be an object: {slide!r}")
    return slides


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Append slides to a .pptx with the same builder the helper CLI uses.",
        epilog="Exit codes: 0 appended, 2 bad input, 3 write failure.",
    )
    parser.add_argument("--input", required=True, help="deck to extend")
    parser.add_argument("--spec", help="JSON file: {slides:[...]} or a list of slides")
    parser.add_argument("--slide", action="append", default=[], help="one slide as inline JSON (repeatable)")
    parser.add_argument("--out", help="output deck (default: <input>-edited.pptx)")
    args = parser.parse_args()

    source = Path(args.input)
    if not source.is_file():
        return _emit({"ok": False, "error": f"no such file: {source}"}, 2)
    output = Path(args.out) if args.out else source.with_name(f"{source.stem}-edited{source.suffix}")

    try:
        slides = _load_slides(args.spec, args.slide)
    except (OSError, ValueError, json.JSONDecodeError) as error:
        return _emit({"ok": False, "error": f"bad spec: {error}"}, 2)
    if not slides:
        return _emit({"ok": False, "error": "nothing to add: pass --spec or --slide"}, 2)

    if not CLI_DIR.is_dir():
        return _emit({"ok": False, "error": f"helper CLI not found at {CLI_DIR}"}, 2)
    sys.path.insert(0, str(CLI_DIR))
    try:
        import office_cli  # noqa: PLC0415 - optional dependency, resolved at run time
    except ImportError as error:
        return _emit({"ok": False, "error": f"cannot import the slide builder: {error}"}, 2)

    from pptx import Presentation  # noqa: PLC0415

    try:
        prs = Presentation(str(source))
    except Exception as error:  # noqa: BLE001
        return _emit({"ok": False, "error": f"cannot open {source}: {error}"}, 2)

    before = len(prs.slides)
    counters = {key: 0 for key in COUNTERS}
    added: list[dict] = []
    for slide_spec in slides:
        try:
            slide = office_cli._pptx_add_slide(prs, slide_spec, counters)  # noqa: SLF001
        except Exception as error:  # noqa: BLE001 - report which slide failed
            return _emit(
                {"ok": False, "error": f"slide {len(added) + 1} rejected: {error}", "added": added},
                2,
            )
        added.append(
            {
                "position": len(prs.slides),
                "layout": str(slide_spec.get("layout", "bullets")),
                "title": slide_spec.get("title"),
                "placeholder_count": len(slide.placeholders),
            }
        )

    try:
        output.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(output))
    except Exception as error:  # noqa: BLE001
        return _emit({"ok": False, "error": f"cannot write {output}: {error}"}, 3)

    return _emit(
        {
            "schema": "add_slides/1",
            "ok": True,
            "input": str(source),
            "output": str(output),
            "added": len(added),
            "before": before,
            "slides": len(prs.slides),
            "counts": counters,
            "details": added,
        }
    )


if __name__ == "__main__":
    sys.exit(main())
